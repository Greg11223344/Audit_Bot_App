from pathlib import Path
from datetime import datetime
import re

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


# =========================
# Helpers
# =========================

def safe_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def safe_filename_text(value) -> str:
    value = safe_text(value)

    invalid_chars = ['\\', '/', ':', '*', '?', '"', '<', '>', '|']

    for char in invalid_chars:
        value = value.replace(char, "_")

    value = value.replace(" ", "_")
    value = re.sub(r"_+", "_", value)

    return value.strip("_")


def get_completed_results(audit_results: dict) -> dict:
    completed = {}

    for key, item in audit_results.items():
        result = safe_text(item.get("result", ""))
        comment = safe_text(item.get("comment", ""))
        image_path = safe_text(item.get("image_path", ""))

        if result or comment or image_path:
            completed[key] = item

    return completed


def calculate_overall_summary(audit_results: dict) -> dict:
    completed_results = get_completed_results(audit_results)

    total_checks = len(audit_results)
    performed = len(completed_results)

    pass_count = 0
    near_fail_count = 0
    fail_count = 0
    na_count = 0

    for item in completed_results.values():
        result = safe_text(item.get("result", ""))

        if result == "Pass":
            pass_count += 1
        elif result == "Near Fail":
            near_fail_count += 1
        elif result == "Fail":
            fail_count += 1
        elif result == "N/A":
            na_count += 1

    counted_checks = performed - na_count

    if counted_checks > 0:
        score = round(((pass_count + near_fail_count * 0.5) / counted_checks) * 100, 1)
    else:
        score = 0

    if total_checks > 0:
        performed_percent = round((performed / total_checks) * 100, 1)
    else:
        performed_percent = 0

    return {
        "total_checks": total_checks,
        "performed": performed,
        "performed_percent": performed_percent,
        "pass": pass_count,
        "near_fail": near_fail_count,
        "fail": fail_count,
        "na": na_count,
        "score": score,
    }


def calculate_module_summaries(audit_results: dict) -> dict:
    module_totals = {}

    for _, item in audit_results.items():
        module = safe_text(item.get("module", "Unknown Module"))

        if module not in module_totals:
            module_totals[module] = {
                "total_checks": 0,
                "performed": 0,
                "pass": 0,
                "near_fail": 0,
                "fail": 0,
                "na": 0,
                "score": 0,
                "performed_percent": 0,
            }

        module_totals[module]["total_checks"] += 1

        result = safe_text(item.get("result", ""))
        comment = safe_text(item.get("comment", ""))
        image_path = safe_text(item.get("image_path", ""))

        is_performed = bool(result or comment or image_path)

        if not is_performed:
            continue

        module_totals[module]["performed"] += 1

        if result == "Pass":
            module_totals[module]["pass"] += 1
        elif result == "Near Fail":
            module_totals[module]["near_fail"] += 1
        elif result == "Fail":
            module_totals[module]["fail"] += 1
        elif result == "N/A":
            module_totals[module]["na"] += 1

    final_summaries = {}

    for module, data in module_totals.items():
        if data["performed"] == 0:
            continue

        counted_checks = data["performed"] - data["na"]

        if counted_checks > 0:
            data["score"] = round(
                ((data["pass"] + data["near_fail"] * 0.5) / counted_checks) * 100,
                1
            )
        else:
            data["score"] = 0

        if data["total_checks"] > 0:
            data["performed_percent"] = round(
                (data["performed"] / data["total_checks"]) * 100,
                1
            )
        else:
            data["performed_percent"] = 0

        final_summaries[module] = data

    return final_summaries


def get_machine_level(score, fail_count):
    if fail_count >= 10:
        return "D - Not Approved"

    if score >= 90:
        return "A - Excellent"
    elif score >= 75:
        return "B - Good"
    elif score >= 60:
        return "C - Needs Improvement"
    else:
        return "D - Not Approved"


def get_report_title(audit_details: dict) -> str:
    audit_type = safe_text(audit_details.get("Audit Type", "Full Audit"))

    if audit_type == "Post Installation Audit":
        return "Post Installation Audit Report"

    return "Machine Audit Report"


def get_report_file_prefix(audit_details: dict) -> str:
    audit_type = safe_text(audit_details.get("Audit Type", "Full Audit"))

    if audit_type == "Post Installation Audit":
        return "Post_Installation_Audit_Report"

    return "Machine_Audit_Report"


def set_cell_text(cell, text, bold=False):
    cell.text = ""

    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(safe_text(text))
    run.bold = bold
    run.font.size = Pt(9)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_details_table(document: Document, audit_details: dict):
    document.add_heading("Audit Details", level=1)

    detail_rows = [
        ("Audit Type", audit_details.get("Audit Type", "Full Audit")),
        ("Checklist File", audit_details.get("Checklist File", "")),
        ("Customer / Site", audit_details.get("Customer / Site", "")),
        ("Machine SN", audit_details.get("Machine SN", "")),
        ("Machine Type", audit_details.get("Machine Type", "")),
        ("Site", audit_details.get("Site", "")),
        ("Region", audit_details.get("Region", "")),
        ("Audit Date", audit_details.get("Audit Date", "")),
        ("Engineer Owner", audit_details.get("Engineer Owner", "")),
        ("Audit Engineer", audit_details.get("Audit Engineer", "")),
    ]

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    set_cell_text(header_cells[0], "Field", bold=True)
    set_cell_text(header_cells[1], "Value", bold=True)

    for field, value in detail_rows:
        row_cells = table.add_row().cells
        set_cell_text(row_cells[0], field, bold=True)
        set_cell_text(row_cells[1], value)


def add_overall_summary(document: Document, summary: dict, level: str):
    document.add_heading("Overall Summary", level=1)

    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
        "Total Checks",
        "Performed",
        "Performed %",
        "Pass",
        "Near Fail",
        "Fail",
        "N/A",
        "Score",
    ]

    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)

    row_cells = table.add_row().cells

    values = [
        summary["total_checks"],
        summary["performed"],
        f"{summary['performed_percent']}%",
        summary["pass"],
        summary["near_fail"],
        summary["fail"],
        summary["na"],
        f"{summary['score']}%",
    ]

    for index, value in enumerate(values):
        set_cell_text(row_cells[index], value)

    p = document.add_paragraph()
    p.add_run("Overall Machine Level: ").bold = True
    p.add_run(level)


def add_module_summary(document: Document, module_summaries: dict):
    document.add_heading("Summary by Module", level=1)

    if not module_summaries:
        document.add_paragraph("No performed checks found.")
        return

    table = document.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
        "Module",
        "Total Checks",
        "Performed",
        "Performed %",
        "Pass",
        "Near Fail",
        "Fail",
        "N/A",
        "Score",
    ]

    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)

    for module, data in module_summaries.items():
        row_cells = table.add_row().cells

        values = [
            module,
            data["total_checks"],
            data["performed"],
            f"{data['performed_percent']}%",
            data["pass"],
            data["near_fail"],
            data["fail"],
            data["na"],
            f"{data['score']}%",
        ]

        for index, value in enumerate(values):
            set_cell_text(row_cells[index], value)


def add_detailed_results(document: Document, audit_results: dict):
    document.add_heading("Detailed Performed Checks", level=1)

    completed_results = get_completed_results(audit_results)

    if not completed_results:
        document.add_paragraph("No performed checks found.")
        return

    current_module = None

    for _, item in completed_results.items():
        module = safe_text(item.get("module", ""))
        sub_module = safe_text(item.get("sub_module", ""))
        check_number = safe_text(item.get("check_number", ""))
        check_description = safe_text(item.get("check_description", ""))
        result = safe_text(item.get("result", ""))
        comment = safe_text(item.get("comment", ""))
        image_path = safe_text(item.get("image_path", ""))

        if module != current_module:
            document.add_heading(module, level=2)
            current_module = module

        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        rows = [
            ("Sub Module", sub_module),
            ("Check Number", check_number),
            ("Check Description", check_description),
            ("Result", result),
            ("Engineer Comment", comment),
        ]

        for label, value in rows:
            row_cells = table.add_row().cells
            set_cell_text(row_cells[0], label, bold=True)
            set_cell_text(row_cells[1], value)

        if image_path:
            image_file = Path(image_path)

            if image_file.exists():
                document.add_paragraph("Attached Image:")
                try:
                    document.add_picture(str(image_file), width=Inches(4.8))
                except Exception:
                    document.add_paragraph(f"Image could not be inserted: {image_path}")
            else:
                document.add_paragraph(f"Image path saved but file was not found: {image_path}")

        document.add_paragraph("")


def generate_word_report(audit_details: dict, audit_results: dict, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    report_title = get_report_title(audit_details)

    title = document.add_heading(report_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    document.add_paragraph("")

    summary = calculate_overall_summary(audit_results)
    level = get_machine_level(summary["score"], summary["fail"])
    module_summaries = calculate_module_summaries(audit_results)

    add_details_table(document, audit_details)
    document.add_paragraph("")

    add_overall_summary(document, summary, level)
    document.add_paragraph("")

    add_module_summary(document, module_summaries)
    document.add_paragraph("")

    add_detailed_results(document, audit_results)

    customer = safe_filename_text(audit_details.get("Customer / Site", "Customer"))
    machine_sn = safe_filename_text(audit_details.get("Machine SN", "Machine"))
    audit_date = safe_filename_text(audit_details.get("Audit Date", datetime.now().strftime("%Y-%m-%d")))
    prefix = get_report_file_prefix(audit_details)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    file_name = f"{prefix}_{customer}_{machine_sn}_{audit_date}_{timestamp}.docx"
    report_path = output_dir / file_name

    document.save(report_path)

    return str(report_path), summary, level